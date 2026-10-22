"""End-to-end API journeys: creator -> checks -> fixes -> submit -> claim-aware review -> receipt,
plus RBAC, the publish gate, policy compile/publish, fact expiry and AI disclosure."""
import json

from app import seed_data as sd
from tests.conftest import as_user


def sample(sid):
    return next(s for s in sd.SAMPLES if s["id"] == sid)


def create(c, s):
    r = c.post("/api/cprs", json={"title": s["title"], "channel": s["channel"], "author_type": s["author_type"], "text": s["text"]})
    assert r.status_code == 200, r.text
    return r.json()["key"]


def detail(c, key):
    r = c.get(f"/api/cprs/{key}")
    assert r.status_code == 200, r.text
    return r.json()


def test_full_journey_with_receipt(client):
    bluefin, meera, rohan = as_user(client, "bluefin"), as_user(client, "meera"), as_user(client, "rohan")
    key = create(bluefin, sample("launch-caption"))
    d = detail(bluefin, key)
    assert d["latest"]["run_status"] == "done"
    assert d["latest"]["summary"]["counts"]["block"] == 4
    assert d["intent"]["must_keep"]

    # gate: cannot submit with open blockers
    r = bluefin.post(f"/api/cprs/{key}/submit")
    assert r.status_code == 409 and "blocking" in r.json()["detail"]

    # accept every verified fix -> v2 is re-checked automatically
    items = [{"finding_id": f["id"]} for f in d["latest"]["findings"] if f["fix_verified"]]
    assert bluefin.post(f"/api/cprs/{key}/fixes", json={"items": items}).status_code == 200
    d = detail(bluefin, key)
    assert d["latest"]["number"] == 2 and d["latest"]["summary"]["counts"]["block"] == 0
    assert "₹349" in d["latest"]["text"]

    assert bluefin.post(f"/api/cprs/{key}/submit").status_code == 200
    d = detail(meera, key)
    assert d["status"] == "in_review" and d["required_roles"] == ["brand_manager", "claims_reviewer"]

    # creators cannot approve
    assert bluefin.post(f"/api/cprs/{key}/review", json={"action": "approve"}).status_code == 403
    assert meera.post(f"/api/cprs/{key}/review", json={"action": "approve", "note": "ok"}).status_code == 200
    assert detail(meera, key)["status"] == "in_review"  # waiting for claims review
    assert rohan.post(f"/api/cprs/{key}/review", json={"action": "approve", "note": "claims ok"}).status_code == 200
    d = detail(rohan, key)
    assert d["status"] == "approved" and d["receipt"]["signature"]

    rec = bluefin.get(f"/api/cprs/{key}/receipt").json()
    assert client.post("/api/receipts/verify", json=rec).json()["valid"] is True
    tampered = json.loads(json.dumps(rec))
    tampered["approved_version"]["number"] = 99
    assert client.post("/api/receipts/verify", json=tampered).json()["valid"] is False

    # approved content is locked
    assert bluefin.post(f"/api/cprs/{key}/versions", json={"text": "changed"}).status_code == 409


def test_dispute_and_override_flow(client):
    asha, meera, rohan = as_user(client, "asha"), as_user(client, "meera"), as_user(client, "rohan")
    key = create(asha, {"title": "Eco post", "channel": "linkedin", "author_type": "human",
                        "text": "Our Aloe Calm Gel is eco-friendly. Find it on our website."})
    d = detail(asha, key)
    f = next(x for x in d["latest"]["findings"] if x["rule_key"] == "C-1")
    assert asha.post(f"/api/cprs/{key}/findings/{f['id']}/dispute", json={"reason": "no"}).status_code == 400  # needs a reason
    assert asha.post(f"/api/cprs/{key}/findings/{f['id']}/dispute", json={"reason": "Refers to the recyclable bottle, see F-11"}).status_code == 200
    assert asha.post(f"/api/cprs/{key}/submit").status_code == 200  # disputed blockers may go to review
    # brand manager cannot override a blocking claim; claims reviewer can, with a reason
    assert meera.post(f"/api/cprs/{key}/findings/{f['id']}/resolve", json={"action": "override", "reason": "fine"}).status_code == 403
    assert meera.post(f"/api/cprs/{key}/review", json={"action": "approve"}).status_code == 409  # unresolved dispute
    assert rohan.post(f"/api/cprs/{key}/findings/{f['id']}/resolve", json={"action": "override", "reason": ""}).status_code == 400
    assert rohan.post(f"/api/cprs/{key}/findings/{f['id']}/resolve", json={"action": "override", "reason": "Bottle claim F-11 covers it"}).status_code == 200
    assert meera.post(f"/api/cprs/{key}/review", json={"action": "approve"}).status_code == 200
    assert rohan.post(f"/api/cprs/{key}/review", json={"action": "approve"}).status_code == 200
    d = detail(meera, key)
    assert d["status"] == "approved"
    assert any(x["action"] == "override" for x in d["receipt"]["decisions"])


def test_creators_only_see_their_own_drafts(client):
    asha, bluefin = as_user(client, "asha"), as_user(client, "bluefin")
    key = create(asha, {"title": "Private", "channel": "email", "author_type": "human", "text": "Hello there, our gel is calm."})
    assert bluefin.get(f"/api/cprs/{key}").status_code == 404
    assert all(c["key"] != key for c in bluefin.get("/api/cprs").json())
    assert client.get(f"/api/cprs/{key}").status_code == 401


def test_policy_import_publish_and_new_rules_fire(client):
    meera, asha = as_user(client, "meera"), as_user(client, "asha")
    doc = (sd.__file__ and open(__import__("pathlib").Path(sd.__file__).resolve().parents[2] / "demo" / "kora-guidelines-2026-update.md", "rb").read())
    assert asha.post("/api/policy/import", files={"file": ("u.md", doc)}).status_code == 403
    r = meera.post("/api/policy/import", files={"file": ("kora-guidelines-2026-update.md", doc)})
    assert r.status_code == 200
    imp = meera.get(f"/api/policy/imports/{r.json()['id']}").json()
    assert imp["status"] == "done" and imp["proposed"] >= 5, imp
    for rule in imp["rules"]:
        assert rule["status"] == "proposed" and rule["source_quote"]
    # proposed rules are not enforced until approved + published
    key = create(asha, {"title": "Spelling", "channel": "instagram", "author_type": "human", "text": "Our gel comes in a lovely green color. Never mention GlowCo."})
    assert not any(f["rule_key"].startswith("T-") and "color" in f["excerpt"] for f in detail(asha, key)["latest"]["findings"])
    for rule in imp["rules"]:
        assert meera.patch(f"/api/policy/rules/{rule['id']}", json={"status": "approved"}).status_code == 200
    pv = meera.post("/api/policy/publish", json={"note": "Social & spelling update"}).json()["number"]
    assert pv == 3
    # submit is blocked until re-checked against the new policy
    r = asha.post(f"/api/cprs/{key}/submit")
    assert r.status_code == 409
    assert asha.post(f"/api/cprs/{key}/run").status_code == 200
    fs = detail(asha, key)["latest"]["findings"]
    ex = {f["excerpt"] for f in fs}
    assert "color" in ex and "GlowCo" in ex
    assert any(f["rule_key"].startswith("P-") and "#KoraNaturals" in f["title"] for f in fs)


def test_fact_expiry_and_recheck(client):
    rohan, asha = as_user(client, "rohan"), as_user(client, "asha")
    key = create(asha, {"title": "Coverage", "channel": "linkedin", "author_type": "human",
                        "text": "Kora Naturals ships to more than 19,000 PIN codes across India."})
    assert detail(asha, key)["latest"]["summary"]["counts"]["block"] == 0
    assert asha.patch("/api/facts/F-19", json={"valid_until": "2026-10-01"}).status_code == 403
    assert rohan.patch("/api/facts/F-19", json={"valid_until": "2026-10-01"}).status_code == 200
    imp = rohan.get("/api/facts/F-19/impact").json()
    assert any(c["key"] == key for c in imp["open"])
    assert rohan.post("/api/facts/F-19/recheck").json()["rechecked"] >= 1
    fs = detail(asha, key)["latest"]["findings"]
    assert any(f["verdict"] == "expired" and f["fact_key"] == "F-19" for f in fs)
    rohan.patch("/api/facts/F-19", json={"clear_valid_until": True})


def test_ai_disclosure_logged_and_confidential_blocks_external(client):
    asha, admin = as_user(client, "asha"), as_user(client, "admin")
    r = asha.post("/api/cprs", json={"title": "Secret", "channel": "email", "author_type": "human", "confidential": True,
                                    "text": "Project Monsoon launches with a new gel."})
    key = r.json()["key"]
    log = asha.get(f"/api/ailog?cpr={key}").json()
    assert log["calls"] and all(c["status"] == "skipped" and c["location"] == "none" for c in log["calls"])
    assert asha.get("/api/ailog").status_code == 403
    assert admin.get("/api/ailog").status_code == 200


def test_upload_docx(client):
    import io

    import docx

    d = docx.Document()
    d.add_paragraph("Aloe Calm Gel launch one-pager")
    d.add_paragraph("Our chemical-free Aloe Gel is the best gel in India.")
    buf = io.BytesIO()
    d.save(buf)
    asha = as_user(client, "asha")
    r = asha.post("/api/cprs/upload", files={"file": ("onepager.docx", buf.getvalue())}, data={"channel": "print"})
    assert r.status_code == 200, r.text
    fs = detail(asha, r.json()["key"])["latest"]["findings"]
    assert {"T-4", "T-2", "C-3"} <= {f["rule_key"] for f in fs}
