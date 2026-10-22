"""Generate binary demo files (DOCX one-pager, PDF guideline update) from text sources.
Run at image build time so no binaries live in git."""
import re
import sys
from pathlib import Path

import docx
import pymupdf as fitz

demo = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parents[2] / "demo")

d = docx.Document()
d.add_heading("Aloe Calm Gel — retail one-pager (draft)", 1)
for p in [
    "Meet Aloe Calm Gel, our chemical-free everyday gel for hot Indian summers.",
    "It is the best soothing gel in India and gives up to 36 hours of hydration.",
    "Made with 96% naturally derived ingredients, in a 100% recyclable PET bottle.",
    "Launch price: ₹299 for 100 ml. Free shipping on orders above ₹499.",
    "Stock up before it's gone. Click here to order.",
]:
    d.add_paragraph(p)
d.save(demo / "kora-onepager.docx")

md = (demo / "kora-guidelines-2026-update.md").read_text(encoding="utf-8")
doc = fitz.open()
page = doc.new_page()
y = 56
for line in md.splitlines():
    if line.startswith(">"):
        continue
    txt = re.sub(r"^#+\s*", "", line).strip().replace("—", "-").replace("“", '"').replace("”", '"').replace("’", "'")
    if not txt:
        y += 6
        continue
    size = 15 if line.startswith("# ") else 12.5 if line.startswith("## ") else 10.5
    rect = fitz.Rect(56, y, 540, y + 400)
    used = page.insert_textbox(rect, txt, fontsize=size, fontname="helv")
    y += (400 - used) + 6 if used >= 0 else 40
    if y > 760:
        page = doc.new_page()
        y = 56
doc.save(demo / "kora-guidelines-2026-update.pdf")
print("demo files written to", demo)
